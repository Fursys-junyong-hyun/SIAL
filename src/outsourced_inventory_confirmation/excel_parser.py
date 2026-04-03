from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import openpyxl
import xlrd

from .models import InventoryRow, SourceFile, VendorPreview

REQUIRED_HEADERS = ("거래처명", "자재코드", "색상")
OPTIONAL_HEADERS = ("자재명",)


@dataclass(slots=True)
class ParsedSource:
    source: SourceFile
    rows: list[InventoryRow]


def _normalize_header(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _find_header_row(rows: Iterable[list[object]]) -> tuple[int, dict[str, int]]:
    for row_index, row in enumerate(rows):
        mapping = {
            _normalize_header(value): idx
            for idx, value in enumerate(row)
            if _normalize_header(value)
        }
        if all(header in mapping for header in REQUIRED_HEADERS):
            return row_index, mapping
    raise ValueError("필수 컬럼을 포함한 헤더 행을 찾을 수 없습니다.")


def _iter_sheet_rows(path: Path, sheet_name: str = "Sheet1") -> list[list[object]]:
    suffix = path.suffix.lower()
    if suffix == ".xls":
        workbook = xlrd.open_workbook(path)
        if sheet_name not in workbook.sheet_names():
            raise ValueError(f"시트 '{sheet_name}'을 찾을 수 없습니다.")
        sheet = workbook.sheet_by_name(sheet_name)
        return [sheet.row_values(idx) for idx in range(sheet.nrows)]

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"시트 '{sheet_name}'을 찾을 수 없습니다.")
    sheet = workbook[sheet_name]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def parse_source_file(source: SourceFile) -> ParsedSource:
    raw_rows = _iter_sheet_rows(source.path)
    header_row_index, mapping = _find_header_row(raw_rows[:10])
    data_rows = raw_rows[header_row_index + 1 :]
    parsed_rows: list[InventoryRow] = []
    seen_keys: set[tuple[str, str, str]] = set()

    for raw_row in data_rows:
        vendor_name = str(raw_row[mapping["거래처명"]]).strip() if mapping["거래처명"] < len(raw_row) else ""
        material_code = str(raw_row[mapping["자재코드"]]).strip() if mapping["자재코드"] < len(raw_row) else ""
        color_code = str(raw_row[mapping["색상"]]).strip() if mapping["색상"] < len(raw_row) else ""
        material_name = ""
        if "자재명" in mapping and mapping["자재명"] < len(raw_row):
            material_name = str(raw_row[mapping["자재명"]]).strip()

        if not vendor_name or not material_code or not color_code:
            continue

        key = (vendor_name, material_code, color_code)
        if key in seen_keys:
            continue
        seen_keys.add(key)

        parsed_rows.append(
            InventoryRow(
                site_name=source.site_name.strip(),
                vendor_name=vendor_name,
                material_code=material_code,
                color_code=color_code,
                material_name=material_name,
            )
        )

    return ParsedSource(source=source, rows=parsed_rows)


def collect_vendor_names(parsed_sources: list[ParsedSource]) -> list[str]:
    vendor_names = sorted({row.vendor_name for source in parsed_sources for row in source.rows})
    return vendor_names


def build_vendor_preview(parsed_sources: list[ParsedSource], vendor_name: str) -> VendorPreview:
    rows = [
        row
        for source in parsed_sources
        for row in source.rows
        if row.vendor_name == vendor_name
    ]
    rows.sort(key=lambda item: (item.site_name, item.material_code, item.color_code))
    return VendorPreview(vendor_name=vendor_name, rows=rows)
