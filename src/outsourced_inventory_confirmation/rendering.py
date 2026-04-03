from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from .models import InventoryRow, RenderedDocument, VendorPreview
from .paths import default_template_path

HEADER_ROWS = (1, 2, 3)
BODY_START_ROW = 4
BODY_END_ROW = 30
TAIL_SUM_ROW = 31
TAIL_SPACER_ROW = 32
TAIL_SIGN_ROW = 33
BODY_CAPACITY = BODY_END_ROW - BODY_START_ROW + 1
COPY_COLUMNS = 8


@dataclass(slots=True)
class LayoutSnapshot:
    row_heights: dict[int, float | None]
    col_widths: dict[str, float | None]


def _title_text(report_date: date) -> str:
    return f"■ 재고자산 확인서 ('{report_date:%y}년 {report_date:%m}월 {report_date:%d}일 기준)"


def safe_vendor_name(vendor_name: str) -> str:
    invalid = '<>:"/\\|?*'
    cleaned = "".join("_" if char in invalid else char for char in vendor_name).strip()
    return cleaned or "업체명없음"


def _copy_cell(source: Cell, target: Cell) -> None:
    if source.data_type == "f":
        target.value = source.value
    else:
        target.value = source.value
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)


def _copy_row_block(template_ws, target_ws, source_start: int, source_end: int, target_start: int) -> None:
    for offset, source_row in enumerate(range(source_start, source_end + 1)):
        target_row = target_start + offset
        source_dim = template_ws.row_dimensions[source_row]
        target_ws.row_dimensions[target_row].height = source_dim.height
        for col in range(1, COPY_COLUMNS + 1):
            _copy_cell(template_ws.cell(source_row, col), target_ws.cell(target_row, col))


def _load_template():
    template_path = default_template_path()
    template_wb = load_workbook(template_path)
    return template_path, template_wb, template_wb[template_wb.sheetnames[0]]


def _layout_snapshot(template_ws) -> LayoutSnapshot:
    col_widths = {col: template_ws.column_dimensions[col].width for col in "ABCDEFGH"}
    row_heights = {row: template_ws.row_dimensions[row].height for row in range(1, TAIL_SIGN_ROW + 1)}
    return LayoutSnapshot(row_heights=row_heights, col_widths=col_widths)


def _chunks(rows: list[InventoryRow], size: int) -> list[list[InventoryRow]]:
    return [rows[index : index + size] for index in range(0, len(rows), size)] or [[]]


def _write_header(target_ws, block_start: int, vendor_name: str, report_date: date) -> None:
    target_ws.cell(block_start + 1, 2).value = _title_text(report_date)
    target_ws.cell(block_start + 2, 2).value = "사업장"
    target_ws.cell(block_start + 2, 3).value = "업체명"
    target_ws.cell(block_start + 2, 4).value = "자재코드"
    target_ws.cell(block_start + 2, 5).value = "색상코드"
    target_ws.cell(block_start + 2, 6).value = "재고수량"
    target_ws.cell(block_start + 2, 7).value = "비고"
    for row in range(block_start + BODY_START_ROW, block_start + TAIL_SIGN_ROW + 1):
        target_ws.cell(row, 3).value = vendor_name if row < block_start + TAIL_SUM_ROW else target_ws.cell(row, 3).value


def _apply_body_row(target_ws, row_index: int, row: InventoryRow) -> None:
    target_ws.cell(row_index, 2).value = row.site_name
    target_ws.cell(row_index, 3).value = row.vendor_name
    target_ws.cell(row_index, 4).value = row.material_code
    target_ws.cell(row_index, 5).value = row.color_code
    target_ws.cell(row_index, 6).value = None
    target_ws.cell(row_index, 7).value = row.material_name or ""


def _apply_tail(target_ws, tail_start: int, vendor_name: str) -> None:
    target_ws.merge_cells(start_row=tail_start, start_column=2, end_row=tail_start, end_column=5)
    target_ws.merge_cells(start_row=tail_start + 2, start_column=2, end_row=tail_start + 2, end_column=7)

    sum_cell = target_ws.cell(tail_start, 2)
    sum_cell.value = "합계"
    sum_cell.alignment = Alignment(horizontal="center", vertical="center")
    sum_cell.font = Font(name=sum_cell.font.name, size=sum_cell.font.size, bold=True)

    qty_cell = target_ws.cell(tail_start, 6)
    qty_cell.value = None
    qty_cell.alignment = Alignment(horizontal="center", vertical="center")

    for col in range(2, 8):
        target_ws.cell(tail_start + 1, col).value = None

    sign_cell = target_ws.cell(tail_start + 2, 2)
    sign_cell.value = f"{vendor_name} (인)"
    sign_cell.alignment = Alignment(horizontal="center", vertical="center")


def render_excel(preview: VendorPreview, report_date: date, output_root: Path) -> RenderedDocument:
    _, _, template_ws = _load_template()
    snapshot = _layout_snapshot(template_ws)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "재고자산확인서 양식"
    worksheet.sheet_view.showGridLines = False

    for col, width in snapshot.col_widths.items():
        worksheet.column_dimensions[col].width = width

    # Excel output stays as one continuous sheet. Only PDF repeats headers by page.
    base_tail_start = max(BODY_START_ROW + len(preview.rows), TAIL_SUM_ROW)
    final_row = base_tail_start + 2

    _copy_row_block(template_ws, worksheet, 1, 3, 1)
    _write_header(worksheet, 1, preview.vendor_name, report_date)

    for row_number in range(BODY_START_ROW, base_tail_start):
        worksheet.row_dimensions[row_number].height = template_ws.row_dimensions[BODY_START_ROW].height
        for col in range(1, COPY_COLUMNS + 1):
            _copy_cell(template_ws.cell(BODY_START_ROW, col), worksheet.cell(row_number, col))
            worksheet.cell(row_number, col).value = None

    for offset, row in enumerate(preview.rows):
        _apply_body_row(worksheet, BODY_START_ROW + offset, row)

    _copy_row_block(template_ws, worksheet, TAIL_SUM_ROW, TAIL_SIGN_ROW, base_tail_start)
    _apply_tail(worksheet, base_tail_start, preview.vendor_name)

    worksheet.print_area = f"A1:H{final_row}"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.orientation = worksheet.ORIENTATION_PORTRAIT
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_margins.left = 0.2
    worksheet.page_margins.right = 0.2
    worksheet.page_margins.top = 0.3
    worksheet.page_margins.bottom = 0.3

    safe_name = safe_vendor_name(preview.vendor_name)
    stamp = date.today().strftime("%y%m%d")
    vendor_dir = output_root / safe_name
    vendor_dir.mkdir(parents=True, exist_ok=True)
    xlsx_path = vendor_dir / f"{safe_name}_재고자산확인서_{stamp}.xlsx"
    workbook.save(xlsx_path)

    pdf_path = vendor_dir / f"{safe_name}_재고자산확인서_{stamp}.pdf"
    render_pdf(preview, report_date, pdf_path)

    return RenderedDocument(
        vendor_name=preview.vendor_name,
        report_date=report_date,
        output_dir=vendor_dir,
        xlsx_path=xlsx_path,
        pdf_path=pdf_path,
    )


def _register_font() -> str:
    candidates = [
        Path("C:/Windows/Fonts/malgun.ttf"),
        Path("C:/Windows/Fonts/malgunbd.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            font_name = "MalgunGothic"
            if font_name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(font_name, str(candidate)))
            return font_name
    return "Helvetica"


def _truncate_text(value: str, limit: int) -> str:
    if len(value) <= limit:
        return value
    return value[: max(0, limit - 1)] + "…"


def render_pdf(preview: VendorPreview, report_date: date, output_path: Path) -> None:
    font_name = _register_font()
    page_width, page_height = A4
    left_margin = 36
    right_margin = 36
    top_margin = 36
    bottom_margin = 36
    usable_width = page_width - left_margin - right_margin

    raw_widths = [48, 84, 110, 56, 56, 170]
    scale = usable_width / sum(raw_widths)
    col_widths = [width * scale for width in raw_widths]

    title_height = 24
    header_height = 24
    data_height = 21
    tail_height = 21

    pdf = canvas.Canvas(str(output_path), pagesize=A4)
    pdf.setTitle(f"{preview.vendor_name} 재고자산확인서")
    pages = _chunks(preview.rows, BODY_CAPACITY)

    for page_index, page_rows in enumerate(pages):
        is_last = page_index == len(pages) - 1
        y = page_height - top_margin
        x_positions = [left_margin]
        for width in col_widths:
            x_positions.append(x_positions[-1] + width)

        pdf.setFont(font_name, 12)
        pdf.drawString(left_margin + 8, y - 16, _title_text(report_date))
        y -= title_height

        headers = ["사업장", "업체명", "자재코드", "색상코드", "재고수량", "비고"]
        pdf.setFont(font_name, 9)
        _draw_grid_row(pdf, x_positions, y, header_height, headers, font_name, alignments=["center"] * 6, bold=True)
        y -= header_height

        for row in page_rows:
            values = [
                _truncate_text(row.site_name, 10),
                _truncate_text(row.vendor_name, 16),
                _truncate_text(row.material_code, 20),
                _truncate_text(row.color_code, 10),
                "",
                _truncate_text(row.material_name, 42),
            ]
            _draw_grid_row(pdf, x_positions, y, data_height, values, font_name, alignments=["center", "center", "center", "center", "center", "left"])
            y -= data_height

        if is_last:
            _draw_merged_row(pdf, left_margin, x_positions[4], y, tail_height, "합계", font_name)
            _draw_cell(pdf, x_positions[4], x_positions[5], y, tail_height, "", font_name, "center")
            _draw_cell(pdf, x_positions[5], x_positions[6], y, tail_height, "", font_name, "left")
            y -= tail_height

            _draw_blank_row(pdf, x_positions, y, tail_height)
            y -= tail_height

            _draw_merged_row(pdf, left_margin, x_positions[6], y, tail_height, f"{preview.vendor_name} (인)", font_name)

        pdf.showPage()

    pdf.save()


def _draw_grid_row(pdf, x_positions, y, height, values, font_name, alignments, bold: bool = False) -> None:
    for index, value in enumerate(values):
        _draw_cell(pdf, x_positions[index], x_positions[index + 1], y, height, value, font_name, alignments[index], bold)


def _draw_blank_row(pdf, x_positions, y, height) -> None:
    for index in range(len(x_positions) - 1):
        _draw_cell(pdf, x_positions[index], x_positions[index + 1], y, height, "", "Helvetica", "left")


def _draw_merged_row(pdf, x_left, x_right, y, height, text, font_name) -> None:
    pdf.setStrokeColor(colors.black)
    pdf.rect(x_left, y - height, x_right - x_left, height)
    pdf.setFont(font_name, 9)
    pdf.drawCentredString((x_left + x_right) / 2, y - height + 6, text)


def _draw_cell(pdf, x_left, x_right, y, height, text, font_name, alignment, bold: bool = False) -> None:
    pdf.setStrokeColor(colors.black)
    pdf.rect(x_left, y - height, x_right - x_left, height)
    pdf.setFont(font_name, 9)
    if alignment == "center":
        pdf.drawCentredString((x_left + x_right) / 2, y - height + 6, str(text))
    else:
        pdf.drawString(x_left + 4, y - height + 6, str(text))
